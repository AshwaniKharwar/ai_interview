import io
import uuid
from unittest.mock import AsyncMock, patch
import pytest
from docx import Document
from httpx import AsyncClient

from apps.auth.models import User
from apps.auth.service import generate_tokens_for_user
from apps.interview.models import InterviewSession, InterviewStatus
from apps.interview.schemas import (
    CandidateProfile,
    ExperienceHighlight,
    ExtractionAndBlueprintResult,
    GapAnalysis,
    InterviewBlueprint,
    InterviewProbingRubric,
    InterviewStage,
    JobProfile,
    PotentialGap,
    StrongAlignment,
    UnverifiedClaim,
)
from apps.interview.services.llm import generate_extraction_and_blueprint, regenerate_stage
from tests.conftest import TestingSessionLocal


@pytest.fixture
async def test_user() -> User:
    async with TestingSessionLocal() as session:
        user = User(
            id=uuid.uuid4(),
            email="test_engineer@example.com",
            google_id="google_sub_12345",
            full_name="Alex Morgan",
            is_active=True,
            role="user",
        )
        session.add(user)
        await session.commit()
        await session.refresh(user)
        return user


@pytest.fixture
async def auth_headers(test_user: User) -> dict[str, str]:
    tokens = generate_tokens_for_user(test_user)
    return {"Authorization": f"Bearer {tokens.access_token}"}


@pytest.fixture
async def second_user() -> User:
    async with TestingSessionLocal() as session:
        user = User(
            id=uuid.uuid4(),
            email="other_candidate@example.com",
            google_id="google_sub_67890",
            full_name="Other User",
            is_active=True,
            role="user",
        )
        session.add(user)
        await session.commit()
        await session.refresh(user)
        return user


@pytest.fixture
async def second_auth_headers(second_user: User) -> dict[str, str]:
    tokens = generate_tokens_for_user(second_user)
    return {"Authorization": f"Bearer {tokens.access_token}"}


def create_sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Alex Morgan", level=1)
    doc.add_paragraph("Senior Distributed Systems Engineer with 6 years experience.")
    doc.add_paragraph("Skills: Go, Python, Kafka, PostgreSQL, Kubernetes.")
    doc.add_paragraph("FinTech Corp - Architected real-time transaction processing system handling 10k TPS.")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def get_mock_extraction_result() -> ExtractionAndBlueprintResult:
    return ExtractionAndBlueprintResult(
        candidate_profile=CandidateProfile(
            name="Alex Morgan",
            seniority_level="Senior",
            skills=["Go", "Kafka", "PostgreSQL"],
            experience=[
                ExperienceHighlight(
                    company="FinTech Corp",
                    title="Senior Software Engineer",
                    duration="2021-Present",
                    key_highlights=["10k TPS transaction system"],
                )
            ],
            education="B.S. in Computer Science",
        ),
        job_profile=JobProfile(
            job_title="Lead Backend Engineer",
            required_skills=["Go", "Kafka", "PostgreSQL"],
            seniority_target="Senior",
            key_responsibilities=["Architect event streaming platforms"],
        ),
        gap_analysis=GapAnalysis(
            match_score=85,
            strong_alignments=[
                StrongAlignment(
                    skill_or_domain="Kafka",
                    evidence="10k TPS transaction system",
                    confidence="HIGH",
                )
            ],
            potential_gaps=[
                PotentialGap(
                    skill_or_domain="Kubernetes Operators",
                    jd_requirement="Nice to have",
                    resume_status="Basic k8s",
                    priority="LOW",
                )
            ],
            unverified_claims=[
                UnverifiedClaim(
                    claim="10k TPS pipeline",
                    verification_goal="Verify partition and consumer strategy",
                    suggested_stage="technical_deep_dive",
                )
            ],
        ),
        blueprint=InterviewBlueprint(
            total_duration_minutes=30,
            difficulty_level="medium",
            interviewer_persona="rigorous_collaborative",
            stages=[
                InterviewStage(
                    stage_id="warmup_background",
                    title="Introduction",
                    duration_minutes=5,
                    min_turns=2,
                    max_turns=4,
                    objective="Intro",
                    transition_criteria="Done intro",
                    starter_questions=["Tell me about yourself."],
                ),
                InterviewStage(
                    stage_id="technical_deep_dive",
                    title="Technical Architecture",
                    duration_minutes=15,
                    min_turns=4,
                    max_turns=8,
                    objective="Technical evaluation",
                    transition_criteria="Done tech",
                    starter_questions=["How did you partition Kafka for 10k TPS?"],
                    probing_rubric=InterviewProbingRubric(
                        expected_signals=["Idempotency", "Consumer lag"],
                        red_flags=["In-memory state"],
                    ),
                ),
                InterviewStage(
                    stage_id="wrap_up",
                    title="Conclusion",
                    duration_minutes=5,
                    min_turns=1,
                    max_turns=2,
                    objective="Wrap up",
                    transition_criteria="Done",
                    starter_questions=["Any questions for us?"],
                ),
            ],
        ),
    )


# --- Strict Fallback Absence Tests ---

@pytest.mark.asyncio
async def test_llm_strictly_fails_without_api_key(monkeypatch):
    monkeypatch.setattr("core.config.settings.GEMINI_API_KEY", "")
    with pytest.raises(ValueError) as exc:
        await generate_extraction_and_blueprint(
            resume_text="Sample resume",
            jd_text="Sample JD",
            job_title="Software Engineer",
        )
    assert "GEMINI_API_KEY is not configured" in str(exc.value)

    with pytest.raises(ValueError) as exc_regen:
        dummy_stage = InterviewStage(
            stage_id="test",
            title="Test",
            duration_minutes=5,
            objective="Obj",
            transition_criteria="Crit",
        )
        await regenerate_stage(dummy_stage, "Software Engineer")
    assert "GEMINI_API_KEY is not configured" in str(exc_regen.value)


# --- Endpoint Tests ---

@pytest.mark.asyncio
async def test_upload_resume_unauthorized(client: AsyncClient):
    docx_bytes = create_sample_docx_bytes()
    files = {"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    resp = await client.post("/api/v1/interviews/upload-resume", files=files)
    assert resp.status_code == 401


@pytest.mark.asyncio
async def test_upload_resume_invalid_format(client: AsyncClient, auth_headers: dict[str, str]):
    files = {"file": ("malicious.exe", b"fake binary", "application/octet-stream")}
    resp = await client.post("/api/v1/interviews/upload-resume", files=files, headers=auth_headers)
    assert resp.status_code == 400
    assert "Unsupported file format" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_upload_resume_empty_file(client: AsyncClient, auth_headers: dict[str, str]):
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    resp = await client.post("/api/v1/interviews/upload-resume", files=files, headers=auth_headers)
    assert resp.status_code == 400
    assert "Uploaded file is empty" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_upload_resume_docx_success(client: AsyncClient, auth_headers: dict[str, str]):
    docx_bytes = create_sample_docx_bytes()
    files = {"file": ("alex_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    resp = await client.post("/api/v1/interviews/upload-resume", files=files, headers=auth_headers)
    assert resp.status_code == 200
    data = resp.json()
    assert "resume_file_url" in data
    assert "Alex Morgan" in data["raw_text"]
    assert "Kafka" in data["raw_text"]
    assert data["char_count"] > 100


@pytest.mark.asyncio
async def test_create_interview_session(client: AsyncClient, auth_headers: dict[str, str]):
    payload = {
        "job_title": "Lead Backend Engineer",
        "company_name": "Acme Payments",
        "job_description_raw": "Looking for Lead Backend Engineer with deep Go, Kafka, and PostgreSQL experience.",
        "resume_raw_text": "Alex Morgan, Senior Engineer. Kafka and PostgreSQL expert.",
        "duration_minutes": 30,
        "difficulty_level": "hard",
        "interviewer_persona": "rigorous_collaborative",
    }
    resp = await client.post("/api/v1/interviews/", json=payload, headers=auth_headers)
    assert resp.status_code == 201
    data = resp.json()
    assert data["job_title"] == "Lead Backend Engineer"
    assert data["company_name"] == "Acme Payments"
    assert data["status"] == "DRAFT"
    assert data["difficulty_level"] == "hard"
    assert "id" in data


@pytest.mark.asyncio
async def test_generate_blueprint_and_polling(client: AsyncClient, auth_headers: dict[str, str]):
    payload = {
        "job_title": "Lead Backend Engineer",
        "company_name": "FinTech Corp",
        "job_description_raw": "Must know Kafka, Go, Distributed Systems, and High TPS architecture.",
        "resume_raw_text": "Alex Morgan, 6 years experience. Architected 10k TPS pipeline with Kafka and PostgreSQL.",
        "duration_minutes": 30,
        "difficulty_level": "medium",
    }
    create_resp = await client.post("/api/v1/interviews/", json=payload, headers=auth_headers)
    session_id = create_resp.json()["id"]

    with patch("apps.interview.service.generate_extraction_and_blueprint", new_callable=AsyncMock) as mock_llm:
        mock_llm.return_value = get_mock_extraction_result()

        # Trigger asynchronous blueprint generation
        gen_resp = await client.post(
            f"/api/v1/interviews/{session_id}/generate-blueprint", headers=auth_headers
        )
        assert gen_resp.status_code == 202

        # Check status
        status_resp = await client.get(f"/api/v1/interviews/{session_id}/status", headers=auth_headers)
        assert status_resp.status_code == 200
        assert status_resp.json()["status"] == "BLUEPRINT_READY"
        assert status_resp.json()["has_blueprint"] is True

        # Fetch full session details
        detail_resp = await client.get(f"/api/v1/interviews/{session_id}", headers=auth_headers)
        assert detail_resp.status_code == 200
        detail = detail_resp.json()

        # Verify Gap Analysis is populated
        assert detail["gap_analysis_json"] is not None
        assert "match_score" in detail["gap_analysis_json"]
        assert len(detail["gap_analysis_json"]["strong_alignments"]) > 0

        # Verify Blueprint is populated with stages, rubrics, and starters
        assert detail["blueprint_json"] is not None
        stages = detail["blueprint_json"]["stages"]
        assert len(stages) >= 3

        tech_stage = next(s for s in stages if s["stage_id"] == "technical_deep_dive")
        assert len(tech_stage["starter_questions"]) > 0
        assert tech_stage["probing_rubric"] is not None
        assert len(tech_stage["probing_rubric"]["expected_signals"]) > 0


@pytest.mark.asyncio
async def test_update_blueprint(client: AsyncClient, auth_headers: dict[str, str]):
    payload = {
        "job_title": "Backend Engineer",
        "job_description_raw": "Python, SQL, Microservices",
        "resume_raw_text": "Experienced Python Engineer",
    }
    create_resp = await client.post("/api/v1/interviews/", json=payload, headers=auth_headers)
    session_id = create_resp.json()["id"]

    with patch("apps.interview.service.generate_extraction_and_blueprint", new_callable=AsyncMock) as mock_llm:
        mock_llm.return_value = get_mock_extraction_result()
        await client.post(f"/api/v1/interviews/{session_id}/generate-blueprint", headers=auth_headers)

    # Fetch existing blueprint
    detail_resp = await client.get(f"/api/v1/interviews/{session_id}", headers=auth_headers)
    blueprint = detail_resp.json()["blueprint_json"]

    # Customize first stage title and duration
    blueprint["stages"][0]["title"] = "Customized Intro Stage"
    blueprint["stages"][0]["duration_minutes"] = 5
    blueprint["total_duration_minutes"] = 45

    put_resp = await client.put(
        f"/api/v1/interviews/{session_id}/blueprint",
        json={"blueprint": blueprint},
        headers=auth_headers,
    )
    assert put_resp.status_code == 200
    updated_data = put_resp.json()
    assert updated_data["duration_minutes"] == 45
    assert updated_data["blueprint_json"]["stages"][0]["title"] == "Customized Intro Stage"


@pytest.mark.asyncio
async def test_regenerate_stage(client: AsyncClient, auth_headers: dict[str, str]):
    payload = {
        "job_title": "Platform Engineer",
        "job_description_raw": "Kubernetes and Go",
        "resume_raw_text": "Platform Engineer experienced with Kubernetes",
    }
    create_resp = await client.post("/api/v1/interviews/", json=payload, headers=auth_headers)
    session_id = create_resp.json()["id"]

    with patch("apps.interview.service.generate_extraction_and_blueprint", new_callable=AsyncMock) as mock_llm, \
         patch("apps.interview.service.regenerate_stage", new_callable=AsyncMock) as mock_regen:
        mock_llm.return_value = get_mock_extraction_result()
        await client.post(f"/api/v1/interviews/{session_id}/generate-blueprint", headers=auth_headers)

        mock_regen.return_value = InterviewStage(
            stage_id="technical_deep_dive",
            title="Technical Architecture",
            duration_minutes=15,
            objective="Deep dive",
            transition_criteria="Evaluated",
            focus_areas=["Kubernetes Operators and custom CRDs"],
            starter_questions=["How do you build custom Kubernetes Operators in Go?"],
        )

        regen_resp = await client.post(
            f"/api/v1/interviews/{session_id}/stages/technical_deep_dive/regenerate",
            json={"user_instructions": "Focus heavily on Kubernetes Operators and custom CRDs"},
            headers=auth_headers,
        )
        assert regen_resp.status_code == 200
        stages = regen_resp.json()["blueprint_json"]["stages"]
        tech_stage = next(s for s in stages if s["stage_id"] == "technical_deep_dive")
        assert any("Kubernetes Operators" in q for q in tech_stage["starter_questions"] + tech_stage["focus_areas"])


@pytest.mark.asyncio
async def test_user_session_isolation(
    client: AsyncClient,
    auth_headers: dict[str, str],
    second_auth_headers: dict[str, str],
):
    payload = {
        "job_title": "Confidential Role",
        "job_description_raw": "Private JD",
        "resume_raw_text": "User 1 Resume",
    }
    create_resp = await client.post("/api/v1/interviews/", json=payload, headers=auth_headers)
    session_id = create_resp.json()["id"]

    # User 2 attempts to view User 1's session -> must be 404
    resp = await client.get(f"/api/v1/interviews/{session_id}", headers=second_auth_headers)
    assert resp.status_code == 404

    # User 2 attempts to update User 1's blueprint -> must be 404
    put_resp = await client.put(
        f"/api/v1/interviews/{session_id}/blueprint",
        json={"blueprint": {"total_duration_minutes": 30, "difficulty_level": "easy", "interviewer_persona": "friendly", "stages": []}},
        headers=second_auth_headers,
    )
    assert put_resp.status_code == 404
