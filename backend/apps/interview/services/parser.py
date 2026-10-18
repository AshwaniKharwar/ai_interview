import io
import re
from typing import Tuple
from pypdf import PdfReader
import pdfplumber
from docx import Document


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB limit


class DocumentParsingError(Exception):
    """Raised when document cannot be read or parsed."""
    pass


def sanitize_extracted_text(text: str) -> str:
    """Sanitizes raw extracted text from resumes or JDs to protect against
    null bytes, control characters, and prompt injection attempts.
    """
    if not text:
        return ""

    # Remove null bytes and non-printable control chars except \n, \r, \t
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", text)

    # Neutralize potential delimiter breakouts
    text = re.sub(r"<\s*/?\s*untrusted_resume_content\s*>", "[del]", text, flags=re.IGNORECASE)
    text = re.sub(r"<\s*/?\s*user_instructions\s*>", "[del]", text, flags=re.IGNORECASE)
    text = re.sub(r"<\s*/?\s*job_description\s*>", "[del]", text, flags=re.IGNORECASE)

    # Normalize excessive newlines (keep max 2 consecutive newlines)
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Normalize excessive whitespace
    text = re.sub(r"[ \t]+", " ", text)

    return text.strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from a PDF document using pypdf, falling back to pdfplumber
    if extraction returns sparse or empty text.
    """
    extracted_text = ""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        page_texts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                page_texts.append(t)
        extracted_text = "\n".join(page_texts).strip()
    except Exception:
        # Fallback to pdfplumber on reader error
        extracted_text = ""

    # If pypdf returned sparse or no text, try pdfplumber (better layout/table recovery)
    if len(extracted_text) < 100:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                plumber_pages = []
                for p in pdf.pages:
                    pt = p.extract_text()
                    if pt:
                        plumber_pages.append(pt)
                plumber_text = "\n".join(plumber_pages).strip()
                if len(plumber_text) > len(extracted_text):
                    extracted_text = plumber_text
        except Exception as e:
            if not extracted_text:
                raise DocumentParsingError(f"Failed to extract text from PDF: {str(e)}") from e

    return sanitize_extracted_text(extracted_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from a Microsoft Word (.docx) document."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_texts.append(row_str)

        all_text = "\n".join(paragraphs + table_texts)
        return sanitize_extracted_text(all_text)
    except Exception as e:
        raise DocumentParsingError(f"Failed to parse Word document: {str(e)}") from e


def parse_resume_document(file_bytes: bytes, filename: str) -> Tuple[str, bool]:
    """Inspects file extension, validates size, and extracts sanitized text.
    Returns:
        (sanitized_text, is_sparse_or_scanned_flag)
    """
    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise DocumentParsingError(f"File size exceeds maximum allowed limit of 10MB.")

    lower_filename = filename.lower()
    if lower_filename.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif lower_filename.endswith((".docx", ".doc")):
        text = extract_text_from_docx(file_bytes)
    else:
        raise DocumentParsingError(f"Unsupported file format for {filename}. Only PDF and DOCX files are supported.")

    is_sparse = len(text) < 150
    return text, is_sparse
